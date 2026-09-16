from __future__ import annotations

import base64
import json
import random
import re
from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Callable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fastapi import APIRouter, Header, HTTPException
from fastapi.responses import StreamingResponse
from document_provenance import document_provenance_headers
from pypdf import PdfReader
from pydantic import BaseModel, Field


TEMPLATE_DIR = Path(__file__).resolve().parent / "templates"

TEMPLATES = [
    {
        "id": "executive-business-plan",
        "name": "Executive Business Plan",
        "description": "A polished investor-ready plan covering strategy, operations, funding, risk, and milestones.",
        "fileName": "executive-business-plan.docx",
        "previewUrl": None,
        "sections": [
            "Executive summary",
            "Company overview",
            "Market opportunity",
            "Products and services",
            "Operations",
            "Marketing and sales",
            "Financial projections",
            "Risk and exit strategy",
        ],
    }
]


class BusinessProfile(BaseModel):
    companyName: str = Field(min_length=2, max_length=180)
    contactName: str = Field(default="", max_length=180)
    contactEmail: str = Field(default="", max_length=180)
    website: str = Field(default="", max_length=240)
    industry: str = Field(min_length=2, max_length=180)
    location: str = Field(default="", max_length=240)
    stage: str = Field(default="startup", max_length=80)
    description: str = Field(min_length=20, max_length=4000)
    productsServices: str = Field(min_length=5, max_length=3000)
    targetCustomers: str = Field(min_length=5, max_length=2000)
    revenueModel: str = Field(default="", max_length=2000)
    team: str = Field(default="", max_length=2000)
    goals: str = Field(default="", max_length=2000)
    fundingNeed: str = Field(default="", max_length=500)
    additionalContext: str = Field(default="", max_length=4000)
    templateId: str | None = Field(default=None, max_length=100)


class Projection(BaseModel):
    year: str
    revenue: str
    assumptions: str = ""


class FundingUse(BaseModel):
    amount: str
    purpose: str


class RiskItem(BaseModel):
    risk: str
    mitigation: str


class LeadershipItem(BaseModel):
    name: str
    role: str


class Milestone(BaseModel):
    period: str
    goal: str


class BusinessPlanDraft(BaseModel):
    templateId: str
    executiveSummary: str
    companyOverview: str
    marketOpportunity: str
    productsServices: list[str] = Field(default_factory=list)
    leadership: list[LeadershipItem] = Field(default_factory=list)
    operationalPlan: list[str] = Field(default_factory=list)
    marketingSalesStrategy: list[str] = Field(default_factory=list)
    competitiveAdvantage: str
    financialProjections: list[Projection] = Field(default_factory=list)
    fundingUses: list[FundingUse] = Field(default_factory=list)
    risks: list[RiskItem] = Field(default_factory=list)
    exitStrategy: list[str] = Field(default_factory=list)
    milestones: list[Milestone] = Field(default_factory=list)
    nextSteps: str


class BusinessPlanDraftRequest(BaseModel):
    profile: BusinessProfile


class BusinessPlanDocumentRequest(BaseModel):
    profile: BusinessProfile
    draft: BusinessPlanDraft


class AgentAttachment(BaseModel):
    name: str = Field(max_length=240)
    mimeType: str = Field(default="application/octet-stream", max_length=120)
    contentBase64: str = Field(max_length=8_000_000)


class DocumentAgentChatRequest(BaseModel):
    agentId: str = Field(max_length=80)
    message: str = Field(default="", max_length=6000)
    answers: dict[str, Any] = Field(default_factory=dict)
    history: list[dict[str, str]] = Field(default_factory=list)
    attachments: list[AgentAttachment] = Field(default_factory=list, max_length=5)
    interventionContext: dict[str, Any] | None = None


class DocumentAgentGenerateRequest(BaseModel):
    agentId: str = Field(max_length=80)
    answers: dict[str, Any] = Field(default_factory=dict)


DOCUMENT_AGENT_FIELDS: dict[str, list[tuple[str, str]]] = {
    "business-plan": [
        ("description", "In your own words, what does the business do on a normal day?"),
        ("goals", "What would you most like the business to achieve in the next one to three years?"),
        ("challenges", "What is the biggest thing making business difficult right now?"),
        ("fundingNeed", "Would money, equipment, people, or training help most—and what would you use it for?"),
    ],
    "strategic-plan": [
        ("vision", "If the business is doing well three years from now, what will be different?"),
        ("strategicPriorities", "What are the two or three most important improvements to work on first?"),
        ("risks", "What is most likely to get in the way of that progress?"),
        ("timeHorizon", "Should this plan cover one year, three years, or five years?"),
    ],
}


def _text_at_least(value: Any, minimum: int, fallback: str) -> str:
    text = str(value or "").strip()
    if len(text) >= minimum:
        return text
    combined = f"{text}. {fallback}".strip(". ")
    return combined[:4000]


def _business_profile_from_answers(answers: dict[str, Any]) -> BusinessProfile:
    company = str(answers.get("companyName") or "The business").strip()
    industry = str(answers.get("industryLocation") or "Small business").strip()
    description = _text_at_least(answers.get("description"), 20, f"{company} operates in {industry} and serves customers through its day-to-day products and services")
    products = _text_at_least(answers.get("productsServices"), 5, description)
    customers = _text_at_least(answers.get("targetCustomers"), 5, "Customers in the business's local and reachable market")
    return BusinessProfile(
        companyName=company,
        industry=industry,
        description=description,
        productsServices=products,
        targetCustomers=customers,
        revenueModel=str(answers.get("revenueModel") or "Revenue is earned from normal sales of the business's products or services."),
        team=str(answers.get("team") or "The owner and current team manage daily operations."),
        goals=str(answers.get("goals") or "Build a stable, sustainable business and serve more customers."),
        fundingNeed=str(answers.get("fundingNeed") or "Funding requirements will be confirmed during implementation planning."),
        additionalContext=str(answers.get("challenges") or ""),
    )

DOCUMENT_REQUESTS = {
    "business-plan": ["Company profile or registration document", "Recent financial statements or estimates", "Product/service information", "Existing market research or customer evidence"],
    "strategic-plan": ["Current business or company profile", "Previous strategy or business plan", "Recent performance or financial reports", "Organisation structure", "Market, programme, or stakeholder research"],
}


def _template_for(template_id: str | None) -> dict[str, Any]:
    if template_id:
        match = next((item for item in TEMPLATES if item["id"] == template_id), None)
        if not match:
            raise HTTPException(status_code=404, detail="Business plan template not found")
        return match
    return random.choice(TEMPLATES)


def _extract_json_object(text: str) -> dict[str, Any] | None:
    fenced = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", text, re.DOTALL)
    candidate = fenced.group(1) if fenced else text
    start, end = candidate.find("{"), candidate.rfind("}")
    if start < 0 or end <= start:
        return None
    try:
        return json.loads(candidate[start : end + 1])
    except json.JSONDecodeError:
        return None


def _split_items(value: str, limit: int = 5) -> list[str]:
    items = [item.strip(" -\t") for item in re.split(r"[\n;]+", value) if item.strip(" -\t")]
    return items[:limit] or [value.strip()]


def _fallback_draft(profile: BusinessProfile, template_id: str) -> BusinessPlanDraft:
    products = _split_items(profile.productsServices)
    goals = _split_items(profile.goals or "Build a repeatable operating model; grow a loyal customer base; improve financial sustainability", 4)
    team = _split_items(profile.team, 3) if profile.team else [f"Founding team - {profile.industry} leadership and delivery"]
    funding = profile.fundingNeed or "To be confirmed after detailed financial modelling"
    location_text = f" in {profile.location}" if profile.location else ""
    revenue_text = profile.revenueModel or "Revenue will be generated through sales of the listed products and services."

    return BusinessPlanDraft(
        templateId=template_id,
        executiveSummary=(
            f"{profile.companyName} is a {profile.stage} business operating in the {profile.industry} sector{location_text}. "
            f"{profile.description.strip()} The business serves {profile.targetCustomers.strip()}. {revenue_text} "
            f"The current funding requirement is {funding}, subject to validation through detailed forecasts."
        ),
        companyOverview=(
            f"{profile.companyName} was established to address a clear need in the {profile.industry} market. "
            f"Its purpose is to deliver practical value to {profile.targetCustomers.strip()} while building a sustainable, "
            f"well-governed enterprise. The business is currently at the {profile.stage} stage."
        ),
        marketOpportunity=(
            f"The initial addressable market consists of {profile.targetCustomers.strip()}{location_text}. "
            "Demand will be validated through customer interviews, pilot sales, repeat-purchase data, and competitor tracking. "
            "The business will refine market sizing and pricing as verified commercial evidence becomes available."
        ),
        productsServices=products,
        leadership=[LeadershipItem(name=item, role="Leadership and delivery") for item in team],
        operationalPlan=[
            "Document the end-to-end delivery process and assign clear owners.",
            "Track quality, turnaround time, customer feedback, and operating costs monthly.",
            "Build reliable supplier and partner relationships with alternatives for critical inputs.",
            "Review legal, safety, data-protection, and industry compliance requirements quarterly.",
        ],
        marketingSalesStrategy=[
            f"Use focused direct outreach to reach {profile.targetCustomers.strip()}.",
            "Build trust through referrals, demonstrations, case studies, and useful educational content.",
            "Track leads, conversion, repeat business, acquisition cost, and customer satisfaction.",
            "Test channels in small campaigns and increase spend only where results are measurable.",
        ],
        competitiveAdvantage=(
            f"{profile.companyName} will compete through focused knowledge of its target customers, responsive service, "
            "consistent delivery, and rapid learning from market evidence. Its advantage must be strengthened through "
            "documented processes, customer proof, and disciplined financial management."
        ),
        financialProjections=[
            Projection(year="Year 1", revenue="To be validated", assumptions="Customer and pricing validation"),
            Projection(year="Year 2", revenue="To be validated", assumptions="Repeat sales and channel growth"),
            Projection(year="Year 3", revenue="To be validated", assumptions="Measured expansion"),
        ],
        fundingUses=[
            FundingUse(amount=funding, purpose="Working capital, market development, operating capability, and growth priorities"),
        ],
        risks=[
            RiskItem(risk="Demand risk", mitigation="Validate willingness to pay through pilots and tracked sales."),
            RiskItem(risk="Cash-flow risk", mitigation="Maintain a rolling cash-flow forecast and control fixed costs."),
            RiskItem(risk="Delivery risk", mitigation="Document processes, quality checks, and backup suppliers."),
            RiskItem(risk="Competitive risk", mitigation="Monitor alternatives and strengthen customer-visible differentiation."),
        ],
        exitStrategy=[
            "Build a profitable and transferable operation with documented systems.",
            "Review owner, investor, acquisition, and succession options as the business matures.",
        ],
        milestones=[Milestone(period=f"Phase {index + 1}", goal=goal) for index, goal in enumerate(goals)],
        nextSteps=(
            "Validate the assumptions in this draft with customer evidence, competitor research, supplier quotations, and a "
            "12-month cash-flow model before presenting the plan to funders or partners."
        ),
    )


def _build_draft(
    profile: BusinessProfile,
    template_id: str,
    call_ai: Callable[..., str],
) -> tuple[BusinessPlanDraft, str]:
    fallback = _fallback_draft(profile, template_id)
    system_prompt = (
        "You are a rigorous SME business-plan agent. Produce concise, investor-ready content grounded only in the supplied "
        "profile. Never invent market statistics, customers, contracts, registrations, historical results, or financial figures. "
        "Use 'to be validated' where evidence is missing. Return one valid JSON object using exactly the keys and shapes in the "
        "provided fallbackDraft. Cover marketing, competition, operations, people, risk, financial assumptions, milestones, and "
        "exit strategy. Keep list sizes suitable for a concise executive business plan."
    )
    try:
        text = call_ai(
            system_prompt,
            {"profile": profile.model_dump(), "fallbackDraft": fallback.model_dump()},
            max_output_tokens=3600,
            response_mime_type="application/json",
        )
        parsed = _extract_json_object(text)
        if not parsed:
            return fallback, "fallback"
        parsed["templateId"] = template_id
        return BusinessPlanDraft.model_validate(parsed), "gemini"
    except Exception:
        return fallback, "fallback"


def _set_cell_text(cell: Any, text: str) -> None:
    cell.text = text.strip()


def _set_cell_items(cell: Any, items: list[str]) -> None:
    cell.text = ""
    if not items:
        return
    first = cell.paragraphs[0]
    first.text = items[0]
    try:
        first.style = "List Bullet"
    except KeyError:
        pass
    for item in items[1:]:
        paragraph = cell.add_paragraph(item)
        try:
            paragraph.style = "List Bullet"
        except KeyError:
            pass


def _populate_document(profile: BusinessProfile, draft: BusinessPlanDraft, template: dict[str, Any]) -> BytesIO:
    template_path = TEMPLATE_DIR / template["fileName"]
    if not template_path.exists():
        raise HTTPException(status_code=500, detail="Configured business plan template is missing")

    document = Document(template_path)
    if len(document.tables) < 18:
        raise HTTPException(status_code=500, detail="Business plan template structure is invalid")

    title_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Title"]
    title_target = next((paragraph for paragraph in title_paragraphs if paragraph.text.strip()), None)
    if title_target:
        title_target.text = f"{profile.companyName} Business Plan"
    for paragraph in title_paragraphs:
        if paragraph is not title_target:
            paragraph.text = ""

    tables = document.tables
    _set_cell_text(tables[0].cell(0, 0), profile.contactName or profile.companyName)
    _set_cell_text(tables[0].cell(0, 1), profile.contactEmail or "Contact details to be confirmed")
    _set_cell_text(tables[1].cell(0, 0), profile.companyName)
    _set_cell_text(tables[1].cell(0, 1), profile.website or profile.location or "Business details to be confirmed")
    _set_cell_text(tables[2].cell(0, 1), draft.executiveSummary)
    _set_cell_text(tables[3].cell(0, 1), draft.companyOverview)
    _set_cell_text(tables[4].cell(0, 1), draft.marketOpportunity)

    product_midpoint = max(1, (len(draft.productsServices) + 1) // 2)
    _set_cell_items(tables[5].cell(0, 1), draft.productsServices[:product_midpoint])
    _set_cell_items(tables[5].cell(0, 2), draft.productsServices[product_midpoint:])
    _set_cell_text(tables[6].cell(0, 1), f"{profile.companyName} is led by a focused team responsible for strategy, delivery, controls, and growth.")

    for index in range(3):
        if index < len(draft.leadership):
            item = draft.leadership[index]
            _set_cell_text(tables[7].cell(1, index), f"{item.name}\n{item.role}")
        else:
            _set_cell_text(tables[7].cell(1, index), "")

    operations_midpoint = max(1, (len(draft.operationalPlan) + 1) // 2)
    _set_cell_text(tables[8].cell(0, 1), f"To scale delivery and strengthen customer outcomes, {profile.companyName} will:")
    _set_cell_items(tables[8].cell(1, 1), draft.operationalPlan[:operations_midpoint])
    _set_cell_items(tables[8].cell(1, 2), draft.operationalPlan[operations_midpoint:])

    marketing_midpoint = max(1, (len(draft.marketingSalesStrategy) + 1) // 2)
    _set_cell_items(tables[9].cell(0, 1), draft.marketingSalesStrategy[:marketing_midpoint])
    _set_cell_items(tables[9].cell(0, 2), draft.marketingSalesStrategy[marketing_midpoint:])
    _set_cell_text(tables[10].cell(0, 1), draft.competitiveAdvantage)

    for index, projection in enumerate(draft.financialProjections[:3]):
        _set_cell_text(tables[12].cell(0, index * 2), f"{projection.year}\n{projection.revenue}\n{projection.assumptions}")

    for row_index in range(4):
        if row_index < len(draft.fundingUses):
            item = draft.fundingUses[row_index]
            _set_cell_text(tables[13].cell(row_index, 0), item.amount)
            _set_cell_text(tables[13].cell(row_index, 1), item.purpose)
        else:
            _set_cell_text(tables[13].cell(row_index, 0), "")
            _set_cell_text(tables[13].cell(row_index, 1), "")

    _set_cell_items(tables[14].cell(0, 1), [f"{item.risk}: {item.mitigation}" for item in draft.risks])
    _set_cell_items(tables[15].cell(0, 1), draft.exitStrategy)

    milestone_cells = [(0, 0), (0, 2), (2, 0), (2, 2)]
    for index, (row_index, col_index) in enumerate(milestone_cells):
        if index < len(draft.milestones):
            item = draft.milestones[index]
            _set_cell_text(tables[16].cell(row_index, col_index), f"{item.period}\n{item.goal}")
        else:
            _set_cell_text(tables[16].cell(row_index, col_index), "")

    _set_cell_text(tables[17].cell(0, 1), draft.nextSteps)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def _attachment_text(attachment: AgentAttachment) -> str:
    try:
        raw = base64.b64decode(attachment.contentBase64, validate=True)
    except Exception:
        return ""
    if len(raw) > 5_000_000:
        return ""
    suffix = Path(attachment.name).suffix.lower()
    try:
        if suffix == ".pdf" or attachment.mimeType == "application/pdf":
            return "\n".join((page.extract_text() or "") for page in PdfReader(BytesIO(raw)).pages)[:20_000]
        if suffix == ".docx" or "wordprocessingml" in attachment.mimeType:
            document = Document(BytesIO(raw))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            return "\n".join(parts)[:20_000]
        return raw.decode("utf-8", errors="ignore")[:20_000]
    except Exception:
        return ""


def _document_agent_chat(payload: DocumentAgentChatRequest, call_ai: Callable[..., str]) -> dict[str, Any]:
    fields = DOCUMENT_AGENT_FIELDS.get(payload.agentId)
    if not fields:
        raise HTTPException(status_code=404, detail="Document agent not found")
    answers = {key: value for key, value in payload.answers.items() if isinstance(value, str) and value.strip()}
    document_context = [
        {"name": item.name, "text": _attachment_text(item)}
        for item in payload.attachments
    ]
    missing_before = [key for key, _ in fields if not str(answers.get(key, "")).strip()]
    current_key = missing_before[0] if missing_before else None

    if payload.message.strip() or document_context:
        system_prompt = (
            "You are a conversational SME document agent. Extract only facts supplied by the user or their attached documents. "
            "Never invent figures, customers, contracts, evidence, or market statistics. Return JSON only with keys answers, acknowledgement, and stepUpdates. "
            "answers may contain only the field keys supplied in fieldGuide. Preserve existing correct answers and update them when the user clarifies. "
            "acknowledgement must be one short natural sentence and must not ask the next question. "
            "For a multi-step intervention, stepUpdates may contain only supplied step ids with status not_started, in_progress, or completed and a short evidence note. "
            "Mark a step completed only when the conversation or document evidence explicitly satisfies its description; otherwise use in_progress."
        )
        try:
            text = call_ai(
                system_prompt,
                {
                    "agentId": payload.agentId,
                    "fieldGuide": {key: question for key, question in fields},
                    "existingAnswers": answers,
                    "currentField": current_key,
                    "message": payload.message,
                    "documents": document_context,
                    "recentHistory": payload.history[-8:],
                    "interventionContext": payload.interventionContext,
                },
                max_output_tokens=1600,
                response_mime_type="application/json",
            )
            parsed = _extract_json_object(text) or {}
            extracted = parsed.get("answers") if isinstance(parsed.get("answers"), dict) else {}
            allowed = {key for key, _ in fields}
            for key, value in extracted.items():
                if key in allowed and isinstance(value, str) and value.strip():
                    answers[key] = value.strip()
            if payload.agentId == "strategic-plan" and current_key and current_key not in answers and payload.message.strip():
                message_text = payload.message.strip()
                normalized_message = message_text.lower()
                word_count = len(re.findall(r"\b\w+\b", message_text))
                action_count = len(re.findall(r"\b(expand|grow|improve|strengthen|develop|build|increase|reduce|enter|launch|enhance|diversify|streamline|retain)\w*\b", normalized_message))
                numbered_items = len(re.findall(r"(?:^|\n)\s*\d+[.)]", message_text))
                uncertainty_text = re.search(r"\b(not sure|don'?t know|still trying|no idea|unsure)\b", normalized_message)
                valid_priorities = current_key == "strategicPriorities" and word_count >= 8 and (action_count >= 2 or numbered_items >= 2)
                valid_general_answer = current_key != "strategicPriorities" and word_count >= 4 and not uncertainty_text
                if valid_priorities or valid_general_answer:
                    answers[current_key] = message_text
            acknowledgement = str(parsed.get("acknowledgement") or "Thanks, I have captured that.").strip()
            supplied_steps = (payload.interventionContext or {}).get("steps", [])
            supplied_ids = {str(step.get("id")) for step in supplied_steps if isinstance(step, dict)}
            raw_step_updates = parsed.get("stepUpdates") if isinstance(parsed.get("stepUpdates"), list) else []
            step_updates = [
                update for update in raw_step_updates
                if isinstance(update, dict)
                and str(update.get("id")) in supplied_ids
                and update.get("status") in {"not_started", "in_progress", "completed"}
            ]
        except Exception:
            if current_key and payload.message.strip():
                answers[current_key] = payload.message.strip()
            acknowledgement = "Thanks, I have captured that."
            step_updates = []
    else:
        acknowledgement = "I will guide you through this one question at a time."
        step_updates = []

    missing = [key for key, _ in fields if not str(answers.get(key, "")).strip()]
    next_question = next((question for key, question in fields if key in missing), "")
    ready = not missing
    if ready:
        reply = f"{acknowledgement} I have enough information to create the first draft. You can ask me to change anything, or create the document now."
    elif next_question.lower() in acknowledgement.lower() or acknowledgement.rstrip().endswith("?"):
        reply = acknowledgement
    else:
        reply = f"{acknowledgement} {next_question}".strip()

    return {
        "ok": True,
        "reply": reply,
        "answers": answers,
        "missingFields": missing,
        "ready": ready,
        "requestedDocuments": DOCUMENT_REQUESTS[payload.agentId],
        "documentsRead": [item["name"] for item in document_context if item["text"]],
        "stepUpdates": step_updates,
    }


def _strategic_plan_content(answers: dict[str, Any], call_ai: Callable[..., str]) -> dict[str, Any]:
    priorities = _split_items(str(answers.get("strategicPriorities") or "Strengthen the core business; Improve customer value; Build organisational capability"), 5)
    fallback = {
        "executiveSummary": f"This strategic plan sets the direction for {answers.get('organizationName', 'the organisation')} over {answers.get('timeHorizon', 'the next three years')}. It converts the stated vision into focused priorities, measurable outcomes, accountable initiatives, and a practical review rhythm.",
        "currentPosition": str(answers.get("currentPosition") or "Current position to be validated."),
        "vision": str(answers.get("vision") or "Vision to be confirmed."),
        "mission": str(answers.get("mission") or "Mission to be confirmed."),
        "values": ["Customer value", "Accountability", "Continuous improvement"],
        "strategicContext": str(answers.get("marketContext") or "External market and operating assumptions must be validated during implementation."),
        "stakeholders": _split_items(str(answers.get("stakeholders") or "Customers; Employees; Partners; Funders"), 6),
        "swot": {
            "strengths": ["Established knowledge and capabilities described in the current-position assessment"],
            "weaknesses": ["Resource and capability gaps require prioritisation"],
            "opportunities": ["Focused execution against the selected strategic priorities"],
            "threats": _split_items(str(answers.get("risks") or "Market and execution uncertainty"), 4),
        },
        "objectives": [
            {
                "priority": priority,
                "objective": f"Deliver measurable progress in {priority.lower()}.",
                "measures": ["Baseline confirmed", "Quarterly target tracked", "Owner reports progress"],
                "initiatives": [f"Define the {priority.lower()} workplan", "Assign an accountable owner", "Review progress quarterly"],
                "owner": "Executive team",
                "timing": str(answers.get("timeHorizon") or "Plan period"),
            }
            for priority in priorities
        ],
        "implementationApproach": "Translate each priority into a funded annual operating plan. Confirm owners, baselines, quarterly targets, dependencies, and decision rights before execution begins.",
        "resources": str(answers.get("resources") or "Resource requirements must be costed and assigned during annual planning."),
        "governance": "The leadership team owns the strategy. Priority owners report monthly, leadership reviews quarterly, and the full plan is refreshed annually when evidence or operating conditions change.",
        "risks": [
            {"risk": item, "response": "Assign an owner, early-warning indicator, mitigation action, and quarterly review."}
            for item in _split_items(str(answers.get("risks") or "Execution capacity; Funding constraints; Market change"), 6)
        ],
    }
    try:
        text = call_ai(
            "Create a rigorous, concise strategic plan from supplied facts only. Do not invent results or statistics. Return valid JSON matching fallback exactly, including arrays and object shapes.",
            {"answers": answers, "fallback": fallback},
            max_output_tokens=4200,
            response_mime_type="application/json",
        )
        parsed = _extract_json_object(text)
        return parsed if isinstance(parsed, dict) else fallback
    except Exception:
        return fallback


def _set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _populate_strategic_document(answers: dict[str, Any], content: dict[str, Any]) -> BytesIO:
    template_path = TEMPLATE_DIR / "strategic-plan-template.docx"
    document = Document(template_path) if template_path.exists() else Document()
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.72)
    section.left_margin = section.right_margin = Inches(0.82)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size, color in [("Title", 30, "1A364A"), ("Heading 1", 18, "D95823"), ("Heading 2", 13, "1A364A")]:
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    header = section.header.paragraphs[0]
    header.text = f"{answers.get('organizationName', 'Organisation')} | Strategic Plan"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(100, 116, 139)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("  |  Confidential strategic working document").font.size = Pt(8)

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_before = Pt(78)
    run = kicker.add_run("STRATEGIC PLAN")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(217, 88, 35)
    title = document.add_paragraph(style="Title")
    title.add_run(str(answers.get("organizationName") or "Organisation"))
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run(str(answers.get("timeHorizon") or "Three-year strategic direction"))
    meta = document.add_paragraph()
    meta.paragraph_format.space_before = Pt(20)
    meta_run = meta.add_run(f"Prepared by Smart Incubation  |  {datetime.now(timezone.utc).strftime('%B %Y')}")
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(100, 116, 139)
    lead = document.add_paragraph()
    lead.paragraph_format.space_before = Pt(64)
    lead_run = lead.add_run("FROM DIRECTION TO EXECUTION")
    lead_run.bold = True
    lead_run.font.size = Pt(10)
    lead_run.font.color.rgb = RGBColor(217, 88, 35)
    document.add_paragraph(str(content.get("executiveSummary") or ""))
    document.add_page_break()

    def add_heading(text: str, level: int = 1) -> None:
        document.add_heading(text, level=level)

    def add_bullets(items: list[Any]) -> None:
        for item in items:
            document.add_paragraph(str(item), style="List Bullet")

    def shade(cell: Any, fill: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:fill"), fill)

    add_heading("Contents")
    for index, entry in enumerate([
        "1. Strategic foundation", "2. Strategic context and SWOT", "3. Strategic priorities and objectives",
        "4. Implementation roadmap", "5. Performance scorecard", "6. Resources and governance",
        "7. Strategic risks", "8. Immediate next steps",
    ], 1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.08)
        number = paragraph.add_run(f"{index}.  ")
        number.bold = True
        paragraph.add_run(entry.split(". ", 1)[-1])
    document.add_page_break()

    add_heading("1. Strategic foundation")
    add_heading("Current position", 2)
    document.add_paragraph(str(content.get("currentPosition") or ""))
    add_heading("Vision", 2)
    document.add_paragraph(str(content.get("vision") or ""))
    add_heading("Mission", 2)
    document.add_paragraph(str(content.get("mission") or ""))
    add_heading("Guiding values", 2)
    add_bullets(content.get("values") or [])

    add_heading("2. Strategic context")
    document.add_paragraph(str(content.get("strategicContext") or ""))
    add_heading("Key stakeholders", 2)
    add_bullets(content.get("stakeholders") or [])
    swot = content.get("swot") if isinstance(content.get("swot"), dict) else {}
    swot_table = document.add_table(rows=2, cols=2)
    swot_table.style = "Table Grid"
    swot_items = [("STRENGTHS", "strengths", "E7F1E8"), ("WEAKNESSES", "weaknesses", "F8E6E3"), ("OPPORTUNITIES", "opportunities", "EEF4DA"), ("THREATS", "threats", "F5E2E2")]
    for cell, (label, key, fill) in zip([cell for row in swot_table.rows for cell in row.cells], swot_items):
        shade(cell, fill)
        heading = cell.paragraphs[0].add_run(label)
        heading.bold = True
        heading.font.color.rgb = RGBColor(26, 54, 74)
        for value in swot.get(key) or []:
            cell.add_paragraph(str(value), style="List Bullet")

    add_heading("3. Strategic priorities and objectives")
    objectives = content.get("objectives") if isinstance(content.get("objectives"), list) else []
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Shading Accent 1"
    table.autofit = False
    widths = [Inches(1.35), Inches(2.35), Inches(1.75), Inches(1.05)]
    for index, (cell, width) in enumerate(zip(table.rows[0].cells, widths)):
        cell.width = width
        cell.text = ["Priority", "Objective and initiatives", "Measures", "Owner / timing"][index]
    _set_repeat_table_header(table.rows[0])
    for item in objectives:
        if not isinstance(item, dict):
            continue
        cells = table.add_row().cells
        cells[0].text = str(item.get("priority") or "")
        cells[1].text = f"{item.get('objective', '')}\n" + "\n".join(str(value) for value in item.get("initiatives") or [])
        cells[2].text = "\n".join(str(value) for value in item.get("measures") or [])
        cells[3].text = f"{item.get('owner', '')}\n{item.get('timing', '')}"
        for cell, width in zip(cells, widths):
            cell.width = width

    add_heading("4. Implementation roadmap")
    roadmap = document.add_table(rows=1, cols=5)
    roadmap.style = "Light Shading Accent 1"
    roadmap.autofit = False
    roadmap_headers = ["Priority", "Key initiative", "Owner", "Timing", "Success measure"]
    for cell, label in zip(roadmap.rows[0].cells, roadmap_headers):
        cell.text = label
        shade(cell, "1A364A")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    _set_repeat_table_header(roadmap.rows[0])
    for item in objectives:
        if not isinstance(item, dict):
            continue
        initiatives = item.get("initiatives") or [item.get("objective") or "Confirm initiative"]
        for initiative in initiatives:
            cells = roadmap.add_row().cells
            values = [item.get("priority", ""), initiative, item.get("owner", ""), item.get("timing", ""), "; ".join(str(value) for value in item.get("measures") or [])]
            for cell, value in zip(cells, values):
                cell.text = str(value)

    add_heading("5. Performance scorecard")
    scorecard = document.add_table(rows=1, cols=4)
    scorecard.style = "Light Shading Accent 1"
    for cell, label in zip(scorecard.rows[0].cells, ["Strategic priority", "Objective", "Measures", "Review rhythm"]):
        cell.text = label
        shade(cell, "D95823")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    _set_repeat_table_header(scorecard.rows[0])
    for item in objectives:
        if isinstance(item, dict):
            cells = scorecard.add_row().cells
            for cell, value in zip(cells, [item.get("priority", ""), item.get("objective", ""), "; ".join(str(value) for value in item.get("measures") or []), "Quarterly"]):
                cell.text = str(value)

    add_heading("6. Resources and governance")
    document.add_paragraph(str(content.get("implementationApproach") or ""))
    add_heading("Resource implications", 2)
    document.add_paragraph(str(content.get("resources") or ""))
    add_heading("Governance and review", 2)
    document.add_paragraph(str(content.get("governance") or ""))

    add_heading("7. Strategic risks")
    risk_table = document.add_table(rows=1, cols=2)
    risk_table.style = "Light Shading Accent 1"
    risk_table.autofit = False
    risk_table.rows[0].cells[0].text = "Risk"
    risk_table.rows[0].cells[1].text = "Response and control"
    _set_repeat_table_header(risk_table.rows[0])
    for item in content.get("risks") or []:
        if not isinstance(item, dict):
            continue
        cells = risk_table.add_row().cells
        cells[0].width, cells[1].width = Inches(2.4), Inches(4.1)
        cells[0].text = str(item.get("risk") or "")
        cells[1].text = str(item.get("response") or "")

    add_heading("8. Immediate next steps")
    add_bullets([
        "Validate assumptions and baselines with the leadership team and key stakeholders.",
        "Confirm objective owners, measures, annual targets, budgets, and dependencies.",
        "Translate this plan into a 12-month operating plan and quarterly review calendar.",
    ])

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def create_business_plan_router(
    call_ai: Callable[..., str],
    require_auth: Callable[[str | None], None],
) -> APIRouter:
    router = APIRouter(prefix="/api", tags=["document-agents"])

    @router.get("/business-plan/templates")
    async def list_templates(authorization: str | None = Header(default=None)) -> dict[str, Any]:
        require_auth(authorization)
        return {"ok": True, "templates": deepcopy(TEMPLATES), "selectionMode": "random"}

    @router.post("/business-plan/draft")
    async def create_draft(
        payload: BusinessPlanDraftRequest,
        authorization: str | None = Header(default=None),
    ) -> dict[str, Any]:
        require_auth(authorization)
        template = _template_for(payload.profile.templateId)
        draft, source = _build_draft(payload.profile, template["id"], call_ai)
        return {
            "ok": True,
            "draft": draft.model_dump(),
            "template": deepcopy(template),
            "source": source,
            "generatedAt": datetime.now(timezone.utc).isoformat(),
        }

    @router.post("/business-plan/document")
    async def create_document(
        payload: BusinessPlanDocumentRequest,
        authorization: str | None = Header(default=None),
    ) -> StreamingResponse:
        require_auth(authorization)
        template = _template_for(payload.draft.templateId)
        output = _populate_document(payload.profile, payload.draft, template)
        safe_name = re.sub(r"[^a-zA-Z0-9_-]+", "-", payload.profile.companyName).strip("-") or "business"
        headers = {"Content-Disposition": f'attachment; filename="{safe_name}-business-plan.docx"', **document_provenance_headers(output.getvalue())}
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers=headers,
        )

    @router.post("/document-agents/chat")
    async def document_agent_chat(
        payload: DocumentAgentChatRequest,
        authorization: str | None = Header(default=None),
    ) -> dict[str, Any]:
        require_auth(authorization)
        return _document_agent_chat(payload, call_ai)

    @router.post("/document-agents/document")
    async def create_agent_document(
        payload: DocumentAgentGenerateRequest,
        authorization: str | None = Header(default=None),
    ) -> StreamingResponse:
        require_auth(authorization)
        fields = DOCUMENT_AGENT_FIELDS.get(payload.agentId)
        if not fields:
            raise HTTPException(status_code=404, detail="Document agent not found")
        missing = [key for key, _ in fields if not str(payload.answers.get(key, "")).strip()]
        if missing:
            raise HTTPException(status_code=400, detail=f"The conversation is incomplete: {', '.join(missing)}")

        if payload.agentId == "business-plan":
            answers = payload.answers
            profile = _business_profile_from_answers(answers)
            template = _template_for(None)
            draft, _ = _build_draft(profile, template["id"], call_ai)
            output = _populate_document(profile, draft, template)
            name = profile.companyName
            suffix = "business-plan"
        else:
            content = _strategic_plan_content(payload.answers, call_ai)
            output = _populate_strategic_document(payload.answers, content)
            name = str(payload.answers.get("organizationName") or "organisation")
            suffix = "strategic-plan"

        safe_name = re.sub(r"[^a-zA-Z0-9_-]+", "-", name).strip("-") or "document"
        provenance = document_provenance_headers(output.getvalue())
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{safe_name}-{suffix}.docx"', **provenance},
        )

    @router.post("/document-agents/preview")
    async def preview_agent_document(
        payload: DocumentAgentGenerateRequest,
        authorization: str | None = Header(default=None),
    ) -> dict[str, Any]:
        require_auth(authorization)
        fields = DOCUMENT_AGENT_FIELDS.get(payload.agentId)
        if not fields:
            raise HTTPException(status_code=404, detail="Document agent not found")
        missing = [key for key, _ in fields if not str(payload.answers.get(key, "")).strip()]
        if missing:
            raise HTTPException(status_code=400, detail=f"A few discovery answers are still needed: {', '.join(missing)}")
        if payload.agentId == "business-plan":
            profile = _business_profile_from_answers(payload.answers)
            template = _template_for(None)
            draft, source = _build_draft(profile, template["id"], call_ai)
            return {"ok": True, "agentId": payload.agentId, "source": source, "preview": draft.model_dump()}
        return {"ok": True, "agentId": payload.agentId, "source": "agent", "preview": _strategic_plan_content(payload.answers, call_ai)}

    return router
