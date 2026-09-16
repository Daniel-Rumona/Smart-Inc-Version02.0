"""Build the reusable Word style template used by the Strategic Plan Agent."""

from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "templates" / "strategic-plan-template.docx"
NAVY = RGBColor(26, 54, 74)
ORANGE = RGBColor(217, 88, 35)
MUTED = RGBColor(100, 116, 139)


def set_font(style, name: str, size: float, color: RGBColor, bold: bool = False) -> None:
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold


doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Inches(0.72), Inches(0.72)
section.left_margin, section.right_margin = Inches(0.82), Inches(0.82)
section.header_distance, section.footer_distance = Inches(0.3), Inches(0.35)

styles = doc.styles
set_font(styles["Normal"], "Aptos", 10.5, RGBColor(38, 46, 56))
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Normal"].paragraph_format.line_spacing = 1.12
set_font(styles["Title"], "Aptos Display", 30, NAVY, True)
styles["Title"].paragraph_format.space_after = Pt(8)
set_font(styles["Subtitle"], "Aptos", 14, MUTED)
set_font(styles["Heading 1"], "Aptos Display", 18, ORANGE, True)
styles["Heading 1"].paragraph_format.space_before = Pt(18)
styles["Heading 1"].paragraph_format.space_after = Pt(8)
set_font(styles["Heading 2"], "Aptos Display", 13, NAVY, True)
styles["Heading 2"].paragraph_format.space_before = Pt(11)
styles["Heading 2"].paragraph_format.space_after = Pt(5)
set_font(styles["Heading 3"], "Aptos", 11, NAVY, True)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = header.add_run("STRATEGIC PLAN")
run.font.name, run.font.size, run.font.bold, run.font.color.rgb = "Aptos", Pt(8), True, MUTED
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

doc.add_paragraph("Strategic Plan Template", style="Title")
doc.add_paragraph("Reusable system template", style="Subtitle")
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
